#!/usr/bin/env python3
"""Self-test for profile_source.py and validate_drivepointified.py on synthetic workbooks.

    python3 test_drivepointify.py -v
"""
from __future__ import annotations

import subprocess
import sys
import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
MON = ["JAN", "FEB", "MAR", "APR", "MAY", "JUN", "JUL", "AUG", "SEP", "OCT", "NOV", "DEC"]
KD, KR = "•⚡ Key Driver", "  ⚡ Key Result"


def month_end(y: int, m: int) -> datetime:
    return datetime(y + (m == 12), m % 12 + 1, 1) - (datetime(2000, 1, 2) - datetime(2000, 1, 1))


def source_workbook(path: Path):
    """A customer-style template: 2027 budget block on top, LY block under it, pasted % and a short bridge."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "WHL"
    ws["A1"] = "2027 Budget"
    ws["A2"] = "Channel"
    for i, m in enumerate(MON):
        ws.cell(2, 3 + i, m)
        ws.cell(19, 3 + i, m)
        ws.cell(18, 3 + i, "Act" if i < 8 else "For")
    ws["O2"], ws["Q2"] = "Total", "2026 FC"
    ly = [1000 + 50 * i for i in range(12)]
    for i in range(12):
        c = get_column_letter(3 + i)
        ws[f"{c}20"] = ly[i]                       # LY gross
        ws[f"{c}3"] = ly[i] * 1.1                  # budget gross = LY × 1.1, pasted
        ws[f"{c}4"] = ly[i] * 1.1 * -0.12          # D&R = gross × -12 %, pasted
        ws[f"{c}5"] = f"={c}3+{c}4"                # net
        ws[f"{c}6"] = 500                          # flat $
        ws[f"{c}7"] = 300 + 10 * i                 # irregular $
        ws[f"{c}8"] = f"=SUM({c}6:{c}7)"           # total opex
    ws["A3"], ws["A4"], ws["A5"], ws["A6"], ws["A7"], ws["A8"] = "Gross", "D&R", "Net Sales", "Displays", "Events", "Total"
    ws["A18"], ws["A20"] = "2026 Forecast (LY)", "Gross"
    ws["B30"], ws["C30"] = "Displays", "=SUM(C6:H6)"
    ws["B31"], ws["C31"] = "Events", "=SUM(C7:H7)"
    ws["B32"], ws["C32"] = "Gross", "=SUM(C3:H3)"
    for r in (6, 7, 3):
        pass
    ws["B33"], ws["C33"] = "Net", "=SUM(C5:H5)"
    ws["C34"] = "=SUM(C30:C33)"
    wb.create_sheet("Plan Settings").sheet_state = "hidden"
    wb.save(path)


def chrome(ws, start: tuple[int, int], months: int):
    ws["C2"], ws["C3"] = "End of Period", "Period Type"
    y, m = start
    for i in range(months):
        yy, mm = y + (m - 1 + i) // 12, (m - 1 + i) % 12 + 1
        ws.cell(2, 11 + i, month_end(yy, mm))
        ws.cell(3, 11 + i, "Actual" if (yy, mm) <= (2026, 8) else "Forecast")


def settings(wb):
    st = wb.create_sheet("Settings")
    for k, h in enumerate(["id", "Setting", "Value", "Description"]):
        st.cell(1, 2 + k, h)
    rows = [("settings.smartmodelSpec", "6.0"), ("settings.modelVersion", "0.1.0"), ("settings.modelName", "t"),
            ("settings.modelType", "model"), ("settings.modelStartDate", datetime(2026, 8, 1)),
            ("settings.historicalStartDate", datetime(2026, 1, 1)), ("settings.companyId", "acme"),
            ("settings.lastDateActuals", datetime(2026, 8, 31))]
    for i, (k, v) in enumerate(rows):
        st.cell(2 + i, 2, k)
        st.cell(2 + i, 4, v)
    wb.create_sheet("Index")


def bad_build(path: Path):
    """What a cell-shift 'conversion' looks like: 2027-only spine, LY stacked below, totals beside, $ drivers."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "WHL"
    chrome(ws, (2027, 1), 12)
    for i in range(12):
        ws.cell(3, 11 + i, "Forecast")
        g = (1000 + 50 * i) * 1.1
        ws.cell(20, 11 + i, g)
        ws.cell(21, 11 + i, g * -0.12)
        ws.cell(22, 11 + i, f"={get_column_letter(11 + i)}20+{get_column_letter(11 + i)}21")
        ws.cell(40, 11 + i, MON[i])                 # stacked LY header
        ws.cell(41, 11 + i, 1000 + 50 * i)
    ws["W20"] = "=SUM(K20:V20)"                     # total column beside the spine
    for r, (a, b, c) in {20: (KD, "whl_gross", "Gross"), 21: (KD, "whl_dr", "D&R"), 22: (KR, "whl_net", "Net Sales")}.items():
        ws.cell(r, 1, a), ws.cell(r, 2, b), ws.cell(r, 3, c)
    ws["C40"] = "2026 Forecast (LY)"
    settings(wb)
    wb.save(path)


def good_build(path: Path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "WHL"
    chrome(ws, (2026, 1), 24)
    rows = {20: (KD, "whl_growth", "Growth factor"), 21: (KR, "whl_gross", "Gross"), 22: (KD, "whl_drPct", "D&R %"),
            23: (KR, "whl_dr", "D&R"), 24: (KR, "whl_net", "Net Sales")}
    for r, (a, b, c) in rows.items():
        ws.cell(r, 1, a), ws.cell(r, 2, b), ws.cell(r, 3, c)
    for i in range(24):
        col, ly = get_column_letter(11 + i), get_column_letter(11 + i - 12) if i >= 12 else None
        if i < 12:
            ws[f"{col}21"] = f"='R - Seed'!{col}5"
            ws[f"{col}22"] = f"=IFERROR({col}23/{col}21,0)"
            ws[f"{col}23"] = f"='R - Seed'!{col}6"
        else:
            ws[f"{col}20"] = 1.1
            ws[f"{col}21"] = f"={ly}21*{col}20"
            ws[f"{col}22"] = -0.12
            ws[f"{col}23"] = f"={col}21*{col}22"
        ws[f"{col}24"] = f"={col}21+{col}23"
    settings(wb)
    wb.save(path)


def run(script: str, *args) -> tuple[int, str]:
    p = subprocess.run([sys.executable, str(HERE / script), *map(str, args)], capture_output=True, text=True)
    return p.returncode, p.stdout + p.stderr


def engine_build(path: Path) -> bool:
    sys.path.insert(0, str(HERE))
    import drivepointify_engine as e
    m = e.Model(company_id="acme", company_name="Acme", model_name="Acme Wholesale 2027 Budget",
                spine_start=(2026, 1), months=24, budget_start=(2027, 1), last_actuals=(2026, 8), model_start=(2026, 8))
    seed = m.seed("R - Acme WHL Seed", source_note="synthetic", flags=["Act"] * 8 + ["For"] * 4)
    seed.add("gross", "Gross", [1000 + 50 * i for i in range(12)])
    seed.add("dr", "D&R", [-(1000 + 50 * i) * 0.1 for i in range(12)])
    t = m.schedule("WHL", name="Wholesale Schedule", template_id="acme-whl-budget", description="synthetic")
    t.section("Revenue", "Gross = same month LY × growth factor; D&R = % of gross.")
    t.row("lyGross", "calc", "Gross, same month LY", bud="{ly}{@gross}")
    t.row("growth", "driver", "Growth factor vs LY (×)", ident="whl_growth", values=[1.1] * 12, fmt=e.FMT_X)
    t.row("gross", "result", "Gross", ident="whl_gross", hist=seed.ref("gross"), bud="{c}{@lyGross}*{c}{@growth}")
    t.row("drPct", "driver", "D&R % of gross", ident="whl_drPct", hist="IFERROR({c}{@dr}/{c}{@gross},0)", values=-0.12,
          fmt=e.FMT_PCT)
    t.row("dr", "result", "D&R", ident="whl_dr", hist=seed.ref("dr"), bud="{c}{@gross}*{c}{@drPct}")
    t.row("net", "result", "Net Sales", ident="whl_net", hist="{c}{@gross}+{c}{@dr}", bud="{c}{@gross}+{c}{@dr}", total=True)
    t.row("end", "result", "Cumulative net", ident="whl_cum", hist="{c}{@net}", bud="{p}{@end}+{c}{@net}")
    m.summary("WHL", [("Revenue", "section", None, None), ("Gross", "sum", "gross", e.FMT_USD),
                      ("Net Sales", "total", "net", e.FMT_USD), ("Net %", "ratio", ("net", "gross"), e.FMT_PCT)])
    _, recalculated = m.save(path)
    return recalculated


class DrivepointifyTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.td = tempfile.TemporaryDirectory()
        d = Path(cls.td.name)
        cls.src, cls.bad, cls.good = d / "src.xlsx", d / "bad.xlsx", d / "good.xlsx"
        source_workbook(cls.src)
        bad_build(cls.bad)
        good_build(cls.good)

    def test_profiler_finds_stacked_block_pasted_pct_and_bridge(self):
        code, out = run("profile_source.py", self.src)
        self.assertEqual(code, 0, out)
        self.assertIn("More than one month block", out)
        self.assertIn("pasted % × base", out)           # D&R = gross × -12 %
        self.assertIn("growth factor", out)             # gross = LY × 1.1
        self.assertIn("Plan Settings", out)
        self.assertIn("Total", out)                     # columns after the months

    def test_validator_rejects_cell_shift_conversion(self):
        code, out = run("validate_drivepointified.py", self.bad)
        self.assertEqual(code, 1, out)
        for needle in ("history must sit IN the time series", "0 Actual month", "second month-header",
                       "right of the spine", "exact % of another row"):
            self.assertIn(needle, out)
        self.assertTrue(any(l.startswith("FAIL") and needle in l for l in out.splitlines()
                            for needle in ("second month-header",)), out)

    def test_engine_build_passes_the_gate(self):
        out_path = Path(self.td.name) / "engine.xlsx"
        recalculated = engine_build(out_path)
        args = [out_path] + ([] if recalculated else ["--allow-uncalculated"])
        code, out = run("validate_drivepointified.py", *args)
        self.assertEqual(code, 0, out)
        if recalculated:
            wb = openpyxl.load_workbook(out_path, data_only=True)
            ws = wb["WHL"]
            gross_row = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 2).value == "whl_gross")
            self.assertAlmostEqual(ws.cell(gross_row, 11 + 12).value, 1000 * 1.1, places=6)   # Jan-27 = Jan-26 × 1.1

    def test_validator_accepts_time_series_build_on_structure(self):
        code, out = run("validate_drivepointified.py", self.good)
        structural = [l for l in out.splitlines() if l.startswith("FAIL") and " WHL " in f" {l[6:15].strip()} "]
        self.assertEqual(structural, [], out)           # protocol/calc FAILs expected: synthetic file is not recalculated


if __name__ == "__main__":
    unittest.main()
