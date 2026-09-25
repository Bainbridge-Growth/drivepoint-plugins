#!/usr/bin/env python3
"""
profile_source.py — read a customer's budget / forecast tab BEFORE drivepointifying it.

Finds the things that turn into rework if you only move cells around:
  * month-header rows and the blocks under them (budget vs last-year "LY" blocks, Act/For flags)
  * hard-keyed rows that are really  % × another row     (pasted values -> the % is the driver)
  * hard-keyed rows that are really  LY × a factor       (growth-factor builds)
  * rows that are exact copies of another row             (link them, don't re-key)
  * formula rows of the form  =<row> + <constant>         (LY + adjustment -> adjustment driver)
  * date blocks that are shifted vs the month header      (e.g. Dec … Nov under Jan … Dec)
  * SUM bridges (1H/2H, YoY) that skip a row of the section they mirror
  * opening balances that do not roll from the prior block's ending balance
  * labels typed in column B instead of A, duplicate labels, stale hidden "Plan Settings" tabs

Usage:
    python3 profile_source.py <file.xlsx> [--sheet NAME] [--out report.md]

Output is a markdown report; paste its findings into the build spec and resolve every line
(build it, or write down why not) before building.
"""
from __future__ import annotations

import argparse
import re
import sys
from dataclasses import dataclass, field
from datetime import date, datetime
from itertools import combinations

import openpyxl
from openpyxl.utils import get_column_letter

MONTHS = ["JAN", "FEB", "MAR", "APR", "MAY", "JUN", "JUL", "AUG", "SEP", "OCT", "NOV", "DEC"]
REL = 1e-9


def is_month_token(v) -> bool:
    if isinstance(v, (datetime, date)):
        return True
    return isinstance(v, str) and v.strip().upper()[:3] in MONTHS and len(v.strip()) <= 9


def num(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


@dataclass
class Block:
    header_row: int
    month_cols: list[int]
    title: str
    flags: list[str] = field(default_factory=list)
    rows: list[int] = field(default_factory=list)


class Tab:
    def __init__(self, path: str, sheet: str | None):
        self.wf = openpyxl.load_workbook(path)
        self.wv = openpyxl.load_workbook(path, data_only=True)
        self.name = sheet or self.wf.sheetnames[0]
        self.f, self.v = self.wf[self.name], self.wv[self.name]
        self.max_row, self.max_column = self.f.max_row, self.f.max_column

    def label(self, r: int) -> str:
        for c in (1, 2, 3):
            x = self.f.cell(r, c).value
            if isinstance(x, str) and x.strip() and not x.startswith("="):
                return x.strip()
        return ""

    def label_col(self, r: int) -> int | None:
        for c in (1, 2, 3):
            x = self.f.cell(r, c).value
            if isinstance(x, str) and x.strip() and not x.startswith("="):
                return c
        return None

    def vals(self, r: int, cols) -> list:
        return [num(self.v.cell(r, c).value) for c in cols]

    def formulas(self, r: int, cols) -> list:
        return [self.f.cell(r, c).value for c in cols]


def find_blocks(t: Tab) -> list[Block]:
    heads = []
    for r in range(1, t.max_row + 1):
        cols = [c for c in range(1, min(t.max_column, 60) + 1) if is_month_token(t.v.cell(r, c).value)]
        # a header is >=12 month tokens in a row, or a row of formulas pointing at one (=C2)
        refs = [c for c in range(1, min(t.max_column, 60) + 1)
                if isinstance(t.f.cell(r, c).value, str) and re.fullmatch(r"=\$?[A-Z]{1,3}\$?\d+", t.f.cell(r, c).value)
                and is_month_token(t.v.cell(r, c).value)]
        cols = sorted(set(cols) | set(refs))
        if len(cols) >= 12:
            heads.append((r, cols[:12]))
    blocks = []
    for i, (r, cols) in enumerate(heads):
        title = t.label(r) or t.label(r - 1)
        flags = []
        for rr in (r - 1, r + 1):
            fl = [str(t.v.cell(rr, c).value or "") for c in cols]
            if sum(x.strip().lower().startswith(("act", "for", "bud", "fc")) for x in fl) >= 6:
                flags = fl
                title = t.label(rr) or title
        end = heads[i + 1][0] if i + 1 < len(heads) else t.max_row + 1
        b = Block(r, cols, title, flags)
        b.rows = [rr for rr in range(r + 1, end) if any(num(t.v.cell(rr, c).value) is not None for c in cols)]
        blocks.append(b)
    return blocks


def row_kind(t: Tab, r: int, cols) -> str:
    fs = t.formulas(r, cols)
    n_f = sum(isinstance(x, str) and x.startswith("=") for x in fs)
    n_c = sum(num(x) is not None for x in fs)
    if n_f == len(cols):
        return "formula"
    if n_c == len(cols):
        return "constant"
    if n_f and n_c:
        return "mixed"
    return "partial"


def const_ratio(a: list, b: list) -> float | None:
    pairs = [(x, y) for x, y in zip(a, b) if x is not None and y not in (None, 0)]
    if len(pairs) < 10 or all(abs(x) < 1e-12 for x, _ in pairs):
        return None
    ks = [x / y for x, y in pairs]
    k0 = ks[0]
    if all(abs(k - k0) <= max(REL, abs(k0) * 1e-7) for k in ks):
        return k0
    return None


def clean_monthly_ratio(a: list, b: list) -> list[float] | None:
    """Per-month ratio that is a 'round' number every month (e.g. 2.8%, 3.2%) -> planner typed a % per month."""
    pairs = [(x, y) for x, y in zip(a, b) if x is not None and y not in (None, 0)]
    if len(pairs) < 10 or all(abs(x) < 1e-12 for x, _ in pairs):
        return None
    ks = [x / y for x, y in pairs]
    if len({round(k, 6) for k in ks}) < 2:
        return None
    if all(abs(k - round(k, 6)) <= 1e-9 for k in ks) and all(abs(k) < 5 for k in ks):
        return ks
    return None


def profile(path: str, sheet: str | None) -> str:
    t = Tab(path, sheet)
    out = [f"# Source profile — `{path.split('/')[-1]}` › `{t.name}`", ""]
    blocks = find_blocks(t)
    others = [s for s in t.wf.sheetnames if s != t.name]
    out.append(f"Sheets: {t.wf.sheetnames}. Used range {t.f.dimensions}.")
    if "Plan Settings" in t.wf.sheetnames:
        out.append("- **Hidden `Plan Settings` tab present** — it carries the SharePoint id of *this* file. Do not copy it into "
                   "the drivepointified workbook; the server writes a fresh one on upload.")
    out += ["", "## Blocks (month-header rows)", ""]
    if not blocks:
        out.append("No 12-month header rows found — inspect manually.")
    for b in blocks:
        cl = f"{get_column_letter(b.month_cols[0])}:{get_column_letter(b.month_cols[-1])}"
        fl = f" · flags {sorted(set(x for x in b.flags if x))}" if b.flags else ""
        first = t.v.cell(b.header_row, b.month_cols[0]).value
        dn = ""
        if isinstance(first, (datetime, date)):
            nxt = t.v.cell(b.header_row, b.month_cols[-1] + 1).value
            dn = (f" · **dated {first:%b %Y} → {t.v.cell(b.header_row, b.month_cols[-1]).value:%b %Y}"
                  + (f" plus {nxt:%b %Y} in the next column" if isinstance(nxt, (datetime, date)) else "")
                  + "** — map by these dates, not by column position")
        out.append(f"- row {b.header_row} `{b.title}` · months {cl} · {len(b.rows)} numeric rows{fl}{dn}")
    if len(blocks) > 1:
        out += ["", "> **More than one month block.** Last-year / actuals blocks stacked under the budget must become the "
                "HISTORY columns of ONE date spine (seed R-tab → same rows), not a second block. Map each LY row to its "
                "budget row below."]
    if not blocks:
        return "\n".join(out)
    history = [b for b in blocks if b.flags or re.search(r"\bLY\b|actual|\bact\b", b.title, re.I)]
    all_series = {}
    for budget in [b for b in blocks if b not in history]:
        cols = budget.month_cols

        # ---- rows of the budget block ------------------------------------------------------------
        out += ["", f"## Budget block rows (`{budget.title}`)", "", "| row | label | kind | finding |", "|---|---|---|---|"]
        candidates = [r for r in budget.rows]
        series = {r: t.vals(r, cols) for r in candidates}
        ly_rows = {}
        for b in blocks:
            if b is budget:
                continue
            for r in b.rows:
                ly_rows[r] = t.vals(r, b.month_cols)
        labels_seen: dict[str, int] = {}
        findings_n = 0
        for r in candidates:
            lab, kind = t.label(r), row_kind(t, r, cols)
            notes = []
            lc = t.label_col(r)
            if lc and lc != 1:
                notes.append(f"label typed in column {get_column_letter(lc)}, not A")
            key = lab.lower()
            if key and key in labels_seen:
                notes.append(f"duplicate label (also row {labels_seen[key]}) — make the friendly name unique")
            labels_seen.setdefault(key, r)
            a = series[r]
            if kind == "constant":
                flat = len({round(x or 0, 6) for x in a}) == 1
                if flat:
                    notes.append("flat $ every month")
                everything = {**{x: series[x] for x in candidates}, **ly_rows}
                refs_r = lambda r2: re.search(rf"(?<![0-9]){r}(?![0-9])", str(t.f.cell(r2, cols[0]).value or ""))  # noqa: E731
                copy_of = next((r2 for r2, s2 in everything.items()
                                if r2 != r and not refs_r(r2) and any(a) and all(x is not None and y is not None and abs(x - y) <= 1e-6 * max(1, abs(x))
                                                              for x, y in zip(a, s2))), None)
                if copy_of and not flat:
                    f2 = t.f.cell(copy_of, (cols if copy_of in series else next(b.month_cols for b in blocks if copy_of in b.rows))[0]).value
                    if isinstance(f2, str) and f2.startswith("="):
                        notes.append(f"**pasted value of the FORMULA in row {copy_of} `{t.label(copy_of)}`** — rebuild that logic here")
                    else:
                        notes.append(f"**identical to row {copy_of} `{t.label(copy_of)}`** — link one to the other (or a % of it)")
                hit = None
                if not flat and not copy_of:
                    for r2 in candidates:
                        if r2 == r or len({round(x or 0, 6) for x in series[r2]}) == 1:
                            continue
                        k = const_ratio(a, series[r2])
                        if k is not None and abs(k) < 1:
                            hit = f"**= row {r2} `{t.label(r2)}` × {k:.6g}** (pasted % × base → make the % the Key Driver)"
                            break
                    if not hit:
                        for r2, r3 in combinations([x for x in candidates if x != r], 2):
                            s = [None if (p is None or q is None) else p + q for p, q in zip(series[r2], series[r3])]
                            k = const_ratio(a, s)
                            if k is not None and abs(k) < 1:
                                hit = f"**= (row {r2} + row {r3}) × {k:.6g}** (pasted % × base → make the % the Key Driver)"
                                break
                    if not hit:
                        for r2, ly in ly_rows.items():
                            if len({round(x or 0, 6) for x in ly}) == 1:
                                continue
                            k = const_ratio(a, ly)
                            if k is not None:
                                hit = f"**= row {r2} `{t.label(r2)}` (other block) × {k:.6g}** (growth factor / share → Key Driver)"
                                break
                    if not hit:
                        pool = {**{x: series[x] for x in candidates}, **ly_rows}
                        pcts = {x: v for x, v in pool.items() if x != r and all(p is not None and 0 < abs(p) <= 1.5 for p in v)}
                        for r2, base in pool.items():
                            if r2 == r or hit or len({round(p or 0, 6) for p in base}) == 1:
                                continue
                            for r3, pv in pcts.items():
                                if r3 == r2 or len({round(p, 6) for p in pv}) == 1:
                                    continue
                                if all(x is not None and y and abs(x - y * q) <= 1e-6 * max(1, abs(x)) for x, y, q in zip(a, base, pv)):
                                    hit = (f"**= row {r2} `{t.label(r2)}` × row {r3} `{t.label(r3)}` (a monthly %)** "
                                           "→ that % row is the Key Driver, this row is a Key Result")
                                    break
                    if not hit:
                        for r2 in candidates:
                            if r2 == r:
                                continue
                            ks = clean_monthly_ratio(a, series[r2])
                            if ks and t.label(r2).lower().startswith(("net", "total", "gross")) and all(abs(k) < 1 for k in ks):
                                hit = (f"**= row {r2} `{t.label(r2)}` × a round monthly % "
                                       f"({', '.join(f'{k:.4%}' for k in ks[:4])} …)** → monthly % Key Driver")
                                break
                if hit:
                    notes.append(hit)
            if kind == "formula":
                fs = t.formulas(r, cols)
                adj = [m for m in (re.fullmatch(r"=\$?([A-Z]{1,3})\$?(\d+)([+-][\d.]+)", str(x)) for x in fs) if m]
                refs = {m.group(1) for m in (re.fullmatch(r"=\$?[A-Z]{1,3}\$?(\d+)(?:[+-][\d.]+)?", str(x)) for x in fs) if m}
                if len(refs) == 1:
                    src_row = int(next(iter(refs)))
                    where = "a last-year row" if src_row in ly_rows else "another row"
                    notes.append(f"links to {where} (row {src_row} `{t.label(src_row)}`)"
                                 + (f" **+ hard-coded {adj[0].group(3)} in {len(adj)} month(s)** → adjustment Key Driver" if adj else ""))
            if kind == "mixed":
                fs = t.formulas(r, cols)
                first = fs[0]
                if num(first) is not None and all(isinstance(x, str) for x in fs[1:]):
                    notes.append("opening value hard-keyed, rest formulas (balance roll-forward)")
                    for b in [x for x in blocks if x is not budget]:
                        for r2 in b.rows:
                            if "end" in t.label(r2).lower() and "end" in lab.lower().replace("beginning", "").replace("begin", "") or (
                                    "begin" in lab.lower() and "end" in t.label(r2).lower()):
                                last = num(t.v.cell(r2, b.month_cols[-1]).value)
                                if last is not None and abs(last - num(first)) > 0.5:
                                    notes.append(f"**opening {num(first):,.2f} ≠ LY ending row {r2} ({last:,.2f}); gap "
                                                 f"{num(first) - last:,.2f}** → show an explicit opening-balance adjustment row and ask the owner")
                                break
                        else:
                            continue
                        break
            # date rows inside the block
            dvals = [t.v.cell(r, c).value for c in cols]
            if all(isinstance(x, (datetime, date)) for x in dvals[:1]) or isinstance(dvals[0], (datetime, date)):
                first = dvals[0]
                nxt = t.v.cell(r, cols[-1] + 1).value
                notes.append(f"**date row starting {first:%b %Y}**"
                             + (f" with a 13th column ({nxt:%b %Y})" if isinstance(nxt, (datetime, date)) else "")
                             + " — this sub-block is dated differently from the header; map by these dates, not by column")
            if notes:
                findings_n += 1
                out.append(f"| {r} | {lab or '—'} | {kind} | {'; '.join(notes)} |")
        out.append("")
        out.append(f"{findings_n} rows with findings.")

    series = {r: t.vals(r, blocks[0].month_cols) for b in blocks for r in b.rows}
    budget, cols = blocks[0], blocks[0].month_cols
    # ---- bridges that skip rows ---------------------------------------------------------------
    out += ["", "## Bridges / roll-ups that skip rows", ""]
    bridge_refs: dict[int, list[int]] = {}
    for r in range(1, t.max_row + 1):
        for c in range(1, min(t.max_column, 60) + 1):
            x = t.f.cell(r, c).value
            m = re.fullmatch(r"=SUM\(\$?[A-Z]{1,3}\$?(\d+):\$?[A-Z]{1,3}\$?(\d+)\)", str(x or ""))
            if m and m.group(1) == m.group(2) and int(m.group(1)) != r:
                bridge_refs.setdefault(r, []).append(int(m.group(1)))
    runs, cur = [], []
    for r in sorted(bridge_refs):
        if cur and r != cur[-1] + 1:
            runs.append(cur)
            cur = []
        cur.append(r)
    if cur:
        runs.append(cur)
    flagged = False
    for run in runs:
        refs = sorted({bridge_refs[r][0] for r in run})
        if len(refs) < 3:
            continue
        gaps = [x for x in range(refs[0], refs[-1] + 1) if x not in refs and x in series and t.label(x)
                and not t.label(x).lower().startswith("total")]
        if gaps:
            flagged = True
            out.append(f"- **Bridge rows {run[0]}–{run[-1]} mirror rows {refs[0]}–{refs[-1]} but skip "
                       + ", ".join(f"row {g} `{t.label(g)}`" for g in gaps)
                       + "** — the rebuilt summary will include them; tell the owner their bridge was short.")
    if not flagged:
        out.append("None found.")

    # ---- columns next to the months ----------------------------------------------------------
    extra = []
    for c in range(cols[-1] + 1, min(t.max_column, cols[-1] + 12) + 1):
        hv = t.v.cell(budget.header_row, c).value
        if hv not in (None, ""):
            extra.append(f"{get_column_letter(c)} `{hv}`")
    out += ["", "## Columns after the 12 months", "",
            ("- " + ", ".join(extra) + " → these become a **Budget Summary** tab (SUMIFS over the spine), not "
             "columns on the schedule.") if extra else "None."]
    out += ["", "## Other sheets", "", ", ".join(others) or "none"]
    return "\n".join(out)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--sheet")
    ap.add_argument("--out")
    a = ap.parse_args()
    rep = profile(a.xlsx, a.sheet)
    if a.out:
        with open(a.out, "w") as fh:
            fh.write(rep + "\n")
    print(rep)
    return 0


if __name__ == "__main__":
    sys.exit(main())
