#!/usr/bin/env python3
"""
drivepointify_engine.py — a small, dependency-light engine for building drivepointified workbooks.

Writes, in one pass: Index, Settings (add-in header literals), one or more schedule tabs on a single
month-end date spine starting at K2, a seed R-tab that feeds history into the same rows, a Budget
Summary report tab, recalculated cached values (if the `formulas` package is available) and the
Drivepoint add-in WebExtension part.

    from drivepointify_engine import Model

    m = Model(company_id="acme", company_name="Acme", model_name="Acme Wholesale 2027 Budget",
              spine_start=(2026, 1), months=24, budget_start=(2027, 1), last_actuals=(2026, 8))
    seed = m.seed("R - Acme WHL Seed", source_note="Customer template, 2026-09-24",
                  flags=["Act"] * 8 + ["For"] * 4)
    seed.add("grossWhl", "Wholesale", [..12 FY2026 values..])

    t = m.schedule("WHL", name="Wholesale Schedule", template_id="acme-whl-budget",
                   description="Wholesale channel P&L")
    t.section("Revenue", "Gross = same month LY × growth factor.")
    t.row("lyGross", "calc", "Wholesale gross, same month LY", bud="{ly}{@gross}")
    t.row("growth", "driver", "Growth factor vs LY (×)", ident="whl_growth", values=[1.1] * 12, fmt=FMT_X)
    t.row("gross", "result", "Wholesale", ident="whl_gross", hist=seed.ref("grossWhl"),
          bud="{c}{@lyGross}*{c}{@growth}")

    m.summary("WHL", [("Revenue", "section", None, None), ("Wholesale", "sum", "gross", FMT_USD)])
    path, recalculated = m.save("WHL_2027_Template_drivepointified.xlsx")

Row kinds:  driver (Key Driver marker + id), result (Key Result marker + id), calc (no marker, no id).
Templates:  {c} this column · {ly} 12 columns back · {p} previous column · {m11} 11 back (trailing-12)
            {@key} row of `key` on this tab · {@Tab.key} row on another tab
Months:     `hist` fills the months before budget_start, `bud` / `values` fill budget_start onward.
            `values` = list (one per budget month) or a scalar; typed inputs get the input style.
"""
from __future__ import annotations

import re
import tempfile
import uuid
import zipfile
from datetime import datetime, timedelta
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Color, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

SPINE_COL = 11                                    # K
KD, KR = "•⚡ Key Driver", "  ⚡ Key Result"       # protocol marker strings
FMT_USD = '#,##0_);(#,##0);"-"_)'
FMT_USD2 = '#,##0.00_);(#,##0.00);"-"_)'
FMT_PCT = '0.0%;(0.0%);"-"'
FMT_PCT2 = '0.00%;(0.00%);"-"'
FMT_X = '0.0000'
FMT_DAYS = '#,##0.0_);(#,##0.0);"-"_)'
FMT_DATE = 'mmm-yy'
ADDIN_ID, ADDIN_VERSION = "wa200007015", "6.0.16.0"


def _fill(hex6: str) -> PatternFill:
    return PatternFill(patternType="solid", fgColor=Color(rgb="FF" + hex6), bgColor=Color(rgb="FF" + hex6))


# protocol palette (smartmodel-protocol: colours & typography)
BAR_BLUE, SECTION_BLUE, GRAY, LIGHT, HDR, SET_HDR = "64B0FF", "63AEFF", "808080", "F2F2F2", "E8E8E8", "2D2D2D"
TAB = {"index": "FFFFFF", "settings": "404040", "schedule": "FFC000", "report": "64B0FF"}
F = {
    "w14": Font(name="Calibri", size=14, color="FFFFFF"), "w11b": Font(name="Calibri", size=11, color="FFFFFF", bold=True),
    "w11": Font(name="Calibri", size=11, color="FFFFFF"), "b14b": Font(name="Calibri", size=14, bold=True),
    "b12b": Font(name="Calibri", size=12, bold=True), "b11": Font(name="Calibri", size=11),
    "b11b": Font(name="Calibri", size=11, bold=True), "gi": Font(name="Calibri", size=11, color=GRAY, italic=True),
    "g9i": Font(name="Calibri", size=9, color=GRAY, italic=True), "mono": Font(name="Menlo", size=10),
    "monowb": Font(name="Menlo", size=10, color="FFFFFF", bold=True), "input": Font(name="Calibri", size=11, color="4472C4"),
    "actual": Font(name="Calibri", size=11, color="ED7D31"), "import": Font(name="Calibri", size=11, color="548235"),
}
BORDER_SECTION = Border(bottom=Side(style="thick", color=SECTION_BLUE))
BORDER_TOTAL = Border(top=Side(style="thin", color="000000"))
CENTER, LEFT, RIGHT = (Alignment(horizontal=h, vertical="center") for h in ("center", "left", "right"))
_TOKEN = re.compile(r"\{@(?:([^.}]+)\.)?([^}]+)\}")


def _month_end(y: int, m: int) -> datetime:
    return datetime(y + (m == 12), m % 12 + 1, 1) - timedelta(days=1)


def _idx(start: tuple[int, int], ym: tuple[int, int]) -> int:
    return (ym[0] - start[0]) * 12 + ym[1] - start[1]


class Model:
    def __init__(self, *, company_id: str, company_name: str, model_name: str, spine_start: tuple[int, int],
                 months: int, budget_start: tuple[int, int], last_actuals: tuple[int, int],
                 model_start: tuple[int, int] | None = None, version: str = "0.1.0", currency: str = "USD",
                 source_file: str = "", snapshot: str = ""):
        if not company_id:
            raise ValueError("company_id is the Drivepoint tenant id — required")
        self.company_id, self.company_name, self.model_name, self.version = company_id, company_name, model_name, version
        self.spine_start, self.months, self.currency = spine_start, months, currency
        self.budget_idx = _idx(spine_start, budget_start)
        self.last_actuals, self.model_start = last_actuals, model_start or last_actuals
        self.source_file, self.snapshot = source_file, snapshot
        self.reg: dict[tuple[str, str], int] = {}
        self.wb = openpyxl.Workbook()
        self.wb.remove(self.wb.active)
        self.schedules: list[Schedule] = []
        self.seeds: list[Seed] = []
        self.templates: list[tuple[str, str, str]] = []
        self._summary = None
        if not 0 < self.budget_idx < months:
            raise ValueError("budget_start must fall inside the spine and after its first month")

    # -- helpers -------------------------------------------------------------------------------
    def col(self, i: int) -> str:
        return get_column_letter(SPINE_COL + i)

    @property
    def last_col(self) -> str:
        return self.col(self.months - 1)

    def render(self, template: str, sheet: str, i: int) -> str:
        def sub(m):
            key = (m.group(1) or sheet, m.group(2))
            if key not in self.reg:
                raise KeyError(f"unknown row {key[0]}.{key[1]} referenced from {sheet}")
            return str(self.reg[key])
        out = _TOKEN.sub(sub, template)
        for tok, off in (("{ly}", 12), ("{m11}", 11), ("{p}", 1)):
            if tok in out:
                if i - off < 0:
                    raise ValueError(f"{tok} used in month {i} of {sheet} (before the spine)")
                out = out.replace(tok, self.col(i - off))
        return out.replace("{c}", self.col(i))

    # -- builders ------------------------------------------------------------------------------
    def seed(self, title: str, *, source_note: str = "", flags: list[str] | None = None) -> "Seed":
        s = Seed(self, title, source_note, flags or [])
        self.seeds.append(s)
        return s

    def schedule(self, title: str, *, name: str, template_id: str, description: str) -> "Schedule":
        s = Schedule(self, title, name, template_id, description)
        self.schedules.append(s)
        self.templates.append((template_id, title, description))
        return s

    def summary(self, sheet: str, rows: list, *, title: str = "Budget Summary", template_id: str | None = None,
                description: str = "", periods: list | None = None):
        self._summary = (sheet, rows, title, template_id or f"{self.company_id}-reports", description, periods)

    def save(self, path: str | Path, *, recalc: bool = True) -> tuple[Path, bool]:
        path = Path(path)
        self._index()
        self._settings()
        for s in self.schedules:
            s._write()
        if self._summary:
            self._write_summary(*self._summary)
        for s in self.seeds:
            s._write()
        order = ["Index", "Settings"] + [s.title for s in self.schedules] + ([self._summary[2]] if self._summary else []) \
            + [s.title for s in self.seeds]
        self.wb._sheets = [self.wb[n] for n in order]
        self.wb.active = 2 if self.schedules else 0
        for ws in self.wb.worksheets:
            ws.sheet_view.tabSelected = ws.title == (self.schedules[0].title if self.schedules else "Index")
        self.wb.calculation = CalcProperties(fullCalcOnLoad=True)
        with tempfile.TemporaryDirectory() as td:
            raw = Path(td) / path.name
            self.wb.save(raw)
            calc = Path(td) / ("calc_" + path.name)
            done = recalc and recalc_workbook(raw, calc)
            inject_webextension(calc if done else raw, path)
        return path, bool(done)

    # -- chrome --------------------------------------------------------------------------------
    def chrome(self, ws, *, name_cell: str, tab: str, description: str = "", spine: bool = True):
        ws.sheet_properties.tabColor = TAB[tab]
        ws.sheet_view.showGridLines = False
        width = SPINE_COL + self.months
        for c in range(1, width):
            ws.cell(1, c).fill = _fill(BAR_BLUE)
            ws.cell(2, c).fill, ws.cell(2, c).font, ws.cell(2, c).alignment = _fill("000000"), F["w11b"], CENTER
            ws.cell(3, c).fill, ws.cell(3, c).font, ws.cell(3, c).alignment = _fill(GRAY), F["w11"], CENTER
            ws.cell(4, c).fill = _fill(LIGHT)
        a1 = ws.cell(1, 1, "≡")
        a1.font, a1.alignment, a1.hyperlink = F["w14"], CENTER, "#Index!A1"
        ws.cell(1, 3, f"={name_cell}").font = F["w14"]
        ws.cell(2, 3, "End of Period").alignment = LEFT
        ws.cell(3, 3, "Period Type").alignment = LEFT
        if spine:
            for i in range(self.months):
                c2 = ws.cell(2, SPINE_COL + i, "=EOMONTH(Settings!$D$7,0)" if i == 0 else f"=EOMONTH({self.col(i - 1)}2,1)")
                c2.number_format = FMT_DATE
                ws.cell(3, SPINE_COL + i, f'=IF({self.col(i)}2<=Settings!$D$14,"Actual","Forecast")')
        ws.cell(4, 1, "•").font = F["input"]
        ws.cell(4, 3, "No Errors").font = F["b11"]
        ws.cell(7, 3, f"={name_cell}").font = F["b14b"]
        for c in range(2, width):
            ws.cell(7, c).border = BORDER_SECTION
        if description:
            ws.cell(8, 3, description).font = F["gi"]

    def _index(self):
        ws = self.wb.create_sheet("Index")
        self.chrome(ws, name_cell="D9", tab="index", spine=False)
        ws.cell(1, 1).hyperlink = None
        ws.cell(2, 3, "SmartModel™ Index").font = F["w11b"]
        ws.cell(2, 3).value, ws.cell(3, 3).value = "SmartModel™ Index", None
        ws.cell(4, 3, "Add-in managed — do not edit manually")
        ws.cell(8, 3, "Templates registered in this workbook").font = F["gi"]
        for c, v in ((2, "metadata___name"), (3, "Name"), (4, self.model_name)):
            ws.cell(9, c, v).font = F["mono"] if c == 2 else F["b11"]
        ws.cell(11, 3, "Templates").font = F["b14b"]
        for k, h in enumerate(["Template ID", "Version", "Sheets", "Skill File", "Imports File"]):
            cell = ws.cell(12, 2 + k, h)
            cell.font, cell.fill = F["b11b"], _fill(HDR)
        rows = list(self.templates) + ([(self._summary[3], self._summary[2], "Budget summary report")] if self._summary else [])
        for n, (tid, sheet, _d) in enumerate(rows):
            r = 13 + n
            ws.cell(r, 2, tid).font = F["mono"]
            ws.cell(r, 3, "1.0.0")
            ws.cell(r, 4, sheet)
            ws.cell(r, 5, "(custom template)").font = F["gi"]
            ws.cell(r, 6, ", ".join(s.title for s in self.seeds) or "—").font = F["gi"]
        ws.freeze_panes = "D5"
        for c, w in zip("ABCDEF", (6, 30, 14, 40, 30, 36)):
            ws.column_dimensions[c].width = w

    def _settings(self):
        ws = self.wb.create_sheet("Settings")
        ws.sheet_properties.tabColor = TAB["settings"]
        ws.sheet_view.showGridLines = False
        for k, h in enumerate(["id", "Setting", "Value", "Description"]):   # exact literals — the add-in matches id / Value
            cell = ws.cell(1, 2 + k, h)
            cell.font, cell.fill = F["monowb"], _fill(SET_HDR)
        y0, m0 = self.spine_start
        rows = [
            ("settings.smartmodelSpec", "Protocol Version", "6.0", "SmartModel Protocol version"),
            ("settings.modelVersion", "Model Version", self.version, "Model version (semver)"),
            ("settings.modelName", "Model Name", self.model_name, "Human-readable model name"),
            ("settings.modelType", "Model Type", "model", '"template" or "model"'),
            ("settings.modelStartDate", "ProForma Start Date", datetime(*self.model_start, 1), "First day of the transition month"),
            ("settings.historicalStartDate", "Historical Start Date", datetime(y0, m0, 1), "Start of the date spine (K2)"),
            ("settings.companyId", "Company ID", self.company_id, "Drivepoint tenant id"),
            ("settings.companyName", "Company Name", self.company_name, "Company name"),
            ("settings.currency", "Currency", self.currency, "Base currency (ISO 4217)"),
            ("settings.author", "Author", "Drivepoint", "Author name"),
            ("settings.authorId", "Author ID", "drivepoint", "Author identifier"),
            ("settings.stores", "Stores", "", "Shopify store names (if any)"),
            ("settings.lastDateActuals", "Last Date Actuals", _month_end(*self.last_actuals), "Last month with booked actuals"),
        ]
        for k, (ident, label, value, desc) in enumerate(rows):
            r = 2 + k
            ws.cell(r, 2, ident).font = F["mono"]
            ws.cell(r, 3, label)
            v = ws.cell(r, 4, value)
            if isinstance(value, datetime):
                v.number_format = "YYYY-MM-DD"
            elif ident in ("settings.smartmodelSpec", "settings.modelVersion"):
                v.number_format = "@"
            ws.cell(r, 5, desc).font = F["gi"]
            if r % 2 == 0:
                for c in range(2, 6):
                    ws.cell(r, c).fill = _fill(LIGHT)
        assert ws["B7"].value == "settings.historicalStartDate" and ws["B14"].value == "settings.lastDateActuals"
        ws.cell(16, 3, "Source & conversion notes").font = F["b11b"]
        for k, (ident, label, value) in enumerate([("settings.sourceFile", "Source File", self.source_file),
                                                   ("settings.sourceSnapshot", "Source Snapshot", self.snapshot)]):
            ws.cell(17 + k, 2, ident).font = F["mono"]
            ws.cell(17 + k, 3, label)
            ws.cell(17 + k, 4, value)
        for c, w in zip("ABCDE", (4, 34, 26, 34, 60)):
            ws.column_dimensions[c].width = w

    def _write_summary(self, sheet, rows, title, template_id, description, periods):
        periods = periods or default_periods(self)
        ws = self.wb.create_sheet(title)
        self.chrome(ws, name_cell="D5", tab="report", description=description, spine=False)
        for k, (ident, label, value) in enumerate([("metadata___name", "Name", title),
                                                   ("metadata___template_id", "Template ID", template_id),
                                                   ("metadata___template_version", "Template Version", "1.0.0")]):
            ws.cell(5 + k, 2, ident).font = F["mono"]
            ws.cell(5 + k, 3, label)
            ws.cell(5 + k, 4, value)
        HDR_R, START, END = 10, 11, 12
        ws.cell(HDR_R, 3, "Period").font = F["b11b"]
        for L, label, start, end in periods:
            h = ws[f"{L}{HDR_R}"]
            h.value, h.font, h.fill, h.alignment = label, F["b11b"], _fill(HDR), RIGHT
            if start:
                for rr, d in ((START, start), (END, end)):
                    c = ws[f"{L}{rr}"]
                    c.value, c.number_format, c.font, c.alignment = d, FMT_DATE, F["g9i"], RIGHT
        q = f"'{sheet}'"
        rng = lambda mr: f"{q}!$K${mr}:${self.last_col}${mr}"  # noqa: E731
        dates = f"{q}!$K$2:${self.last_col}$2"
        out: dict[str, int] = {}
        r = 14
        for label, agg, arg, fmt in rows:
            if agg == "section":
                ws.cell(r, 3, label).font = F["b12b"]
                for c in range(2, 20):
                    ws.cell(r, c).border = BORDER_SECTION
                r += 1
                continue
            if agg == "blank":
                r += 1
                continue
            bold = agg == "total"
            ws.cell(r, 3, label).font = F["b11b"] if bold else F["b11"]
            if agg != "ratio":
                out[arg] = r
                mr = self.reg[(sheet, arg)]
            for L, _l, start, _e in periods:
                cell = ws[f"{L}{r}"]
                cell.number_format = fmt
                if start is None:
                    continue
                if agg in ("sum", "total"):
                    cell.value = f'=SUMIFS({rng(mr)},{dates},">="&{L}${START},{dates},"<="&{L}${END})'
                elif agg == "end":
                    cell.value = f"=SUMIFS({rng(mr)},{dates},{L}${END})"
                elif agg == "begin":
                    cell.value = f"=SUMIFS({rng(mr)},{dates},EOMONTH({L}${START},0))"
                else:
                    num, den = arg
                    cell.value = f"=IFERROR({L}{out[num]}/{L}{out[den]},0)"
                cell.font = Font(name="Calibri", size=11, color="548235", bold=bold)
                if bold:
                    cell.border = BORDER_TOTAL
            var_cols = [L for L, _l, s, _e in periods if s is None]
            if len(var_cols) >= 2:
                a, b = [L for L, _l, s, _e in periods if s][:2]
                g, h = ws[f"{var_cols[0]}{r}"], ws[f"{var_cols[1]}{r}"]
                g.value = f"={b}{r}-{a}{r}"
                if agg != "ratio":
                    h.value, h.number_format = f"=IFERROR({var_cols[0]}{r}/ABS({a}{r}),0)", FMT_PCT
            r += 1
        ws.freeze_panes = "E13"
        for c, w in (("A", 6), ("B", 30), ("C", 44), ("D", 3)):
            ws.column_dimensions[c].width = w
        for L, *_ in periods:
            ws.column_dimensions[L].width = 14
        self.summary_rows = out


def default_periods(m: Model) -> list:
    """FY history vs FY budget, variance, budget quarters and halves (columns E…P)."""
    by = m.spine_start[0] + m.budget_idx // 12 if m.spine_start[1] == 1 else None
    by = by or (m.spine_start[0] + (m.spine_start[1] - 1 + m.budget_idx) // 12)
    q = lambda y, a, b: (datetime(y, a, 1), _month_end(y, b))  # noqa: E731
    return [("E", f"FY{by - 1}", *q(by - 1, 1, 12)), ("F", f"FY{by} Budget", *q(by, 1, 12)),
            ("G", f"{by % 100} v {(by - 1) % 100} $", None, None), ("H", f"{by % 100} v {(by - 1) % 100} %", None, None),
            ("J", f"Q1 {by}", *q(by, 1, 3)), ("K", f"Q2 {by}", *q(by, 4, 6)), ("L", f"Q3 {by}", *q(by, 7, 9)),
            ("M", f"Q4 {by}", *q(by, 10, 12)), ("O", f"1H {by}", *q(by, 1, 6)), ("P", f"2H {by}", *q(by, 7, 12))]


class Seed:
    """R-tab holding the customer's history on the model spine; schedule rows read it with SUMIFS on the id."""

    def __init__(self, model: Model, title: str, note: str, flags: list[str]):
        if not re.match(r"^R\s*-\s*", title):
            raise ValueError("seed tab titles must start with 'R - ' (plan sync ignores R-tabs as roll tabs)")
        self.m, self.title, self.note, self.flags = model, title, note, flags
        self.rows: list[tuple[str, str, list]] = []

    def add(self, key: str, label: str, values: list, *, start: tuple[int, int] | None = None):
        """values start at `start` (default: spine start) — e.g. a single Dec value for a TTM block."""
        off = _idx(self.m.spine_start, start) if start else 0
        self.rows.append((key, label, [None] * off + list(values)))

    def ref(self, key: str) -> str:
        return f"SUMIFS('{self.title}'!{{c}}:{{c}},'{self.title}'!$B:$B,\"seed___{key}\")"

    def _write(self):
        m, ws = self.m, self.m.wb.create_sheet(self.title)
        ws.sheet_view.showGridLines = False
        ws.cell(1, 1, "≡").font = F["b14b"]
        ws.cell(1, 3, self.title).font = F["b14b"]
        ws.cell(2, 3, "Line item").font = F["b11b"]
        for i in range(m.months):
            c2 = ws.cell(2, SPINE_COL + i, "=EOMONTH(Settings!$D$7,0)" if i == 0 else f"=EOMONTH({m.col(i - 1)}2,1)")
            c2.number_format, c2.font, c2.fill, c2.alignment = FMT_DATE, F["w11b"], _fill("000000"), CENTER
            c3 = ws.cell(3, SPINE_COL + i, self.flags[i] if i < len(self.flags) else None)
            c3.font, c3.fill, c3.alignment = F["w11"], _fill(GRAY), CENTER
        ws.cell(4, 3, self.note).font = F["g9i"]
        for k, (key, label, vals) in enumerate(self.rows):
            r = 6 + k
            ws.cell(r, 2, f"seed___{key}").font = F["mono"]
            ws.cell(r, 3, label)
            for i, v in enumerate(vals[:m.months]):
                if v is not None:
                    cell = ws.cell(r, SPINE_COL + i, v)
                    cell.number_format = FMT_USD
        ws.freeze_panes = "K4"
        for c, w in zip("ABC", (4, 30, 36)):
            ws.column_dimensions[c].width = w


class Schedule:
    def __init__(self, model: Model, title: str, name: str, template_id: str, description: str):
        self.m, self.title = model, title
        self.ws = model.wb.create_sheet(title)
        model.chrome(self.ws, name_cell="D9", tab="schedule", description=description)
        meta = [("metadata___name", "Name", name), ("metadata___template_id", "Template ID", template_id),
                ("metadata___template_version", "Template Version", "1.0.0"),
                ("metadata___description", "Description", description), ("metadata___type", "Type", "Schedule"),
                ("metadata___grain", "Grain", "monthly"), ("metadata___framework", "Framework", "drivepointify-v1"),
                ("metadata___docsLink", "Documentation", "https://docs.drivepoint.io")]
        for k, (ident, label, value) in enumerate(meta):
            self.ws.cell(9 + k, 2, ident).font = F["mono"]
            self.ws.cell(9 + k, 3, label)
            self.ws.cell(9 + k, 4, value).font = F["gi"] if ident == "metadata___description" else F["b11"]
        self.r = 19
        self.specs: list[dict] = []

    def section(self, title: str, note: str | None = None):
        self.specs.append(dict(kind="section", title=title, note=note, row=self.r))
        self.r += 3 if note else 2

    def blank(self, n: int = 1):
        self.r += n

    def row(self, key: str | None, kind: str, label: str, *, ident: str | None = None, hist: str | None = None,
            bud: str | None = None, values=None, fmt: str = FMT_USD, total: bool = False, bold: bool = False):
        if kind not in ("driver", "result", "calc"):
            raise ValueError("kind is driver | result | calc")
        if kind in ("driver", "result") and not ident:
            raise ValueError(f"{label}: Key Drivers / Results need an id for column B")
        if kind == "driver" and values is None:
            raise ValueError(f"{label}: a Key Driver's budget months are inputs — pass values=")
        if kind == "result" and bud is None:
            raise ValueError(f"{label}: a Key Result's budget months are formulas — pass bud=")
        self.specs.append(dict(kind=kind, key=key, label=label, ident=ident, hist=hist, bud=bud, values=values,
                               fmt=fmt, total=total, bold=bold, row=self.r))
        if key:
            self.m.reg[(self.title, key)] = self.r
        self.r += 1

    def _write(self):
        m, ws = self.m, self.ws
        width = SPINE_COL + m.months
        n_bud = m.months - m.budget_idx
        for s in self.specs:
            r = s["row"]
            if s["kind"] == "section":
                ws.cell(r, 3, s["title"]).font = F["b12b"]
                for c in range(2, width):
                    ws.cell(r, c).border = BORDER_SECTION
                if s["note"]:
                    ws.cell(r + 1, 3, s["note"]).font = F["gi"]
                continue
            if s["kind"] != "calc":
                a = ws.cell(r, 1, KD if s["kind"] == "driver" else KR)
                a.font, a.alignment = F["b11"], LEFT
                ws.cell(r, 2, s["ident"]).font = F["mono"]
            ws.cell(r, 3, s["label"]).font = F["b11b"] if (s["total"] or s["bold"]) else F["b11"]
            vals = s["values"]
            if vals is not None and not isinstance(vals, (list, tuple)):
                vals = [vals] * n_bud
            if vals is not None and len(vals) != n_bud:
                raise ValueError(f"{self.title}.{s['label']}: {len(vals)} budget values, spine has {n_bud}")
            for i in range(m.months):
                cell = ws.cell(r, SPINE_COL + i)
                cell.number_format = s["fmt"]
                if i < m.budget_idx:
                    if s["hist"]:
                        cell.value = "=" + m.render(s["hist"], self.title, i)
                        cell.font = F["import"] if "'R -" in s["hist"] or "'R-" in s["hist"] else F["b11"]
                elif vals is not None:
                    v = vals[i - m.budget_idx]
                    cell.value = 0 if v is None else v
                    cell.font, cell.fill = F["input"], _fill(LIGHT)
                elif s["bud"]:
                    cell.value = "=" + m.render(s["bud"], self.title, i)
                    cell.font = F["b11"]
                if s["total"]:
                    cell.border = BORDER_TOTAL
                if s["total"] or s["bold"]:
                    cell.font = Font(name="Calibri", size=11, color=cell.font.color, bold=True)
        ws.freeze_panes = "K5"
        ws.column_dimensions["A"].width = 15
        ws.column_dimensions["B"].width = 34
        ws.column_dimensions["B"].outline_level = 1
        ws.column_dimensions["C"].width = 46
        ws.column_dimensions["D"].width = 13
        for c in "EFGHIJ":
            ws.column_dimensions[c].width = 3
        for i in range(m.months):
            ws.column_dimensions[m.col(i)].width = 12


# ------------------------------------------------------------------------------ recalc + add-in --
def recalc_workbook(src: Path, dst: Path) -> bool:
    """Compute every formula with the `formulas` package and write the results as cached values.

    Returns False (and leaves `dst` unwritten) when `formulas` is not installed — then open the file in
    Excel, recalculate (Ctrl+Alt+F9) and save before uploading: the add-in reads uncached cells as NaN.
    """
    try:
        import formulas  # noqa: F401
    except ImportError:
        return False
    xl = formulas.ExcelModel().loads(str(src)).finish()
    sol = xl.calculate()
    pat = re.compile(r"^'\[" + re.escape(src.name) + r"\](.+)'!([A-Z]+[0-9]+)$")
    values: dict[tuple[str, str], object] = {}
    for k, v in sol.items():
        mt = pat.match(str(k))
        if not mt:
            continue
        val = getattr(v, "value", v)
        try:
            val = val[0][0]
        except Exception:  # noqa: BLE001
            pass
        if hasattr(val, "item"):
            val = val.item()
        if isinstance(val, (bool, int, float, str)):
            values[(mt.group(1).upper(), mt.group(2))] = val
    esc = lambda t: t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")  # noqa: E731
    with zipfile.ZipFile(src) as zin:
        wbxml = zin.read("xl/workbook.xml").decode()
        rels = zin.read("xl/_rels/workbook.xml.rels").decode()
        rid = {}
        for mt in re.finditer(r"<Relationship\b[^>]*>", rels):
            tag = mt.group(0)
            i, t = re.search(r'Id="([^"]+)"', tag), re.search(r'Target="([^"]+)"', tag)
            if i and t:
                rid[i.group(1)] = t.group(1).lstrip("/")
        files = {}
        for mt in re.finditer(r'<sheet [^>]*name="([^"]+)"[^>]*r:id="([^"]+)"', wbxml):
            tgt = rid[mt.group(2)]
            files[tgt if tgt.startswith("xl/") else "xl/" + tgt] = mt.group(1).replace("&amp;", "&").upper()
        cell_re = re.compile(r'<c r="([A-Z]+[0-9]+)"([^>]*)><f>(.*?)</f><v ?/>(</c>)', re.S)
        n = 0
        with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                sheet = files.get(item.filename)
                if sheet:
                    def repl(mt):
                        nonlocal n
                        ref, attrs, f, close = mt.groups()
                        val = values.get((sheet, ref))
                        attrs = re.sub(r'\s+t="[^"]*"', "", attrs)
                        if isinstance(val, bool):
                            n += 1
                            return f'<c r="{ref}"{attrs} t="b"><f>{f}</f><v>{int(val)}</v>{close}'
                        if isinstance(val, (int, float)):
                            n += 1
                            return f'<c r="{ref}"{attrs}><f>{f}</f><v>{repr(float(val))}</v>{close}'
                        if isinstance(val, str):
                            n += 1
                            kind = "e" if val.startswith("#") else "str"
                            return f'<c r="{ref}"{attrs} t="{kind}"><f>{f}</f><v>{esc(val)}</v>{close}'
                        return mt.group(0)
                    data = cell_re.sub(repl, data.decode()).encode()
                zout.writestr(item, data)
    return n > 0


def inject_webextension(src: Path, dst: Path):
    """Add the Drivepoint add-in WebExtension parts (openpyxl and LibreOffice drop them on save)."""
    guid = str(uuid.uuid4()).upper()
    webext = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
              f'<we:webextension xmlns:we="http://schemas.microsoft.com/office/webextensions/webextension/2010/11" id="{{{guid}}}">'
              f'<we:reference id="{ADDIN_ID}" version="{ADDIN_VERSION}" store="en-US" storeType="OMEX"/>'
              f'<we:alternateReferences><we:reference id="{ADDIN_ID}" version="{ADDIN_VERSION}" store="" storeType="OMEX"/>'
              f'</we:alternateReferences><we:properties/><we:bindings/><we:snapshot/></we:webextension>')
    panes = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
             '<wetp:taskpanes xmlns:wetp="http://schemas.microsoft.com/office/webextensions/taskpanes/2010/11">'
             '<wetp:taskpane dockstate="right" visibility="0" width="350" row="0">'
             '<wetp:webextensionref xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="rId1"/>'
             '</wetp:taskpane></wetp:taskpanes>')
    pane_rels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
                 '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                 '<Relationship Id="rId1" Type="http://schemas.microsoft.com/office/2011/relationships/webextension" '
                 'Target="webextension1.xml"/></Relationships>')
    root_rel = ('<Relationship Id="rId_dp1" Type="http://schemas.microsoft.com/office/2011/relationships/webextensiontaskpanes" '
                'Target="xl/webextensions/taskpanes.xml"/>')
    ct = ('<Override PartName="/xl/webextensions/webextension1.xml" ContentType="application/vnd.ms-office.webextension+xml"/>'
          '<Override PartName="/xl/webextensions/taskpanes.xml" ContentType="application/vnd.ms-office.webextensiontaskpanes+xml"/>')
    src, dst = Path(src), Path(dst)
    tmp = dst.with_suffix(".tmp.xlsx") if src.resolve() == dst.resolve() else dst
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            if item.filename.startswith("xl/webextensions/"):
                continue
            data = zin.read(item.filename)
            if item.filename == "_rels/.rels" and b"webextensiontaskpanes" not in data:
                data = data.replace(b"</Relationships>", root_rel.encode() + b"</Relationships>")
            elif item.filename == "[Content_Types].xml" and b"webextension+xml" not in data:
                data = data.replace(b"</Types>", ct.encode() + b"</Types>")
            zout.writestr(item, data)
        zout.writestr("xl/webextensions/webextension1.xml", webext)
        zout.writestr("xl/webextensions/taskpanes.xml", panes)
        zout.writestr("xl/webextensions/_rels/taskpanes.xml.rels", pane_rels)
    if tmp != dst:
        tmp.replace(dst)


if __name__ == "__main__":
    import sys
    if len(sys.argv) == 3 and sys.argv[1] == "--inject-addin":
        inject_webextension(Path(sys.argv[2]), Path(sys.argv[2]))
        print(f"add-in WebExtension injected into {sys.argv[2]}")
    else:
        print(__doc__)
