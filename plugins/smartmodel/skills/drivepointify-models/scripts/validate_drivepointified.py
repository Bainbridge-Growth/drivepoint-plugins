#!/usr/bin/env python3
"""
validate_drivepointified.py — the hand-off gate for a drivepointified workbook.

Run it on the BUILT (recalculated) workbook. It fails on exactly the things a reviewer would
otherwise have to catch by eye:

  SPINE    one contiguous month-end date spine from K2, nothing to the right of it, row 3 is only
           Actual/Forecast (Actual first), and the actuals boundary matches Settings lastDateActuals
  HISTORY  history lives IN the spine — no Actual-less spine when Settings declares history, no second
           month-header / "Act/For" block stacked under the budget, no Total/LY/variance columns next
           to the spine (those belong on a summary tab)
  MARKERS  column A holds only Key Driver / Key Result markers, column B only ids; every marked row has
           a friendly name in C, ids and names are unique
  MEANING  Key Drivers are inputs in every forecast month (a formula there means it is really a result);
           Key Results are formulas (a hard-keyed result is really a driver); a $ Key Driver that is an
           exact % of another row is a pasted "% × base" — the % must be the driver; hard-keyed
           forecast rows without a marker are listed
  CALC     no error cells, every formula has a cached value (the add-in reads NaN otherwise)
  PROTOCOL Settings header id|Setting|Value|Description, smartmodelSpec 6.0, companyId set, real dates,
           Index tab, WebExtension part, no copied "Plan Settings" tab
  TIES     optional --ties JSON: model rows vs source ranges, value by value

Usage:
    python3 validate_drivepointified.py BUILT.xlsx [--ties ties.json] [--strict]

ties.json:
    [{"model": "WHL!whl_netSales", "cols": "W:AH", "source": "path/to/source.xlsx", "range": "WHL!C19:N19"}, ...]
    "model" may be SHEET!<id in column B> or SHEET!<row number>; add "tol": 0.01 to loosen.

Exit code 1 if any FAIL (or any WARN with --strict).
"""
from __future__ import annotations

import argparse
import json
import re
import sys
import zipfile
from datetime import date, datetime, timedelta

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

SPINE_START = 11                     # column K
KD_TOKEN, KR_TOKEN = "Key Driver", "Key Result"
IGNORE_TABS = re.compile(r"^(Settings|Index|Plan Settings|Home|Key Drivers and Results|Drivepoint Agent Log)$|^[RF]\s*-\s*")
MONTHS = ("JAN", "FEB", "MAR", "APR", "MAY", "JUN", "JUL", "AUG", "SEP", "OCT", "NOV", "DEC")
ERRORS = ("#REF!", "#DIV/0!", "#NAME?", "#VALUE!", "#N/A", "#NUM!", "#NULL!", "Err:")
EPOCH = datetime(1899, 12, 30)


class Report:
    def __init__(self):
        self.lines: list[tuple[str, str, str]] = []

    def add(self, level: str, group: str, msg: str):
        self.lines.append((level, group, msg))

    def ok(self, cond: bool, group: str, msg: str, level_if_bad: str = "FAIL"):
        self.add("PASS" if cond else level_if_bad, group, msg)

    def counts(self):
        return {k: sum(1 for x in self.lines if x[0] == k) for k in ("PASS", "WARN", "FAIL")}

    def render(self) -> str:
        return "\n".join(f"{lvl:<5} {grp:<9} {msg}" for lvl, grp, msg in self.lines)


def as_date(v):
    if isinstance(v, datetime):
        return v
    if isinstance(v, date):
        return datetime(v.year, v.month, v.day)
    if isinstance(v, (int, float)) and not isinstance(v, bool) and v > 20000:
        return EPOCH + timedelta(days=float(v))
    return None


def month_end(d: datetime) -> bool:
    return (d + timedelta(days=1)).day == 1


def is_formula(v) -> bool:
    return isinstance(v, str) and v.startswith("=")


def is_num(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def settings_map(wb) -> dict:
    out = {}
    if "Settings" not in wb.sheetnames:
        return out
    ws = wb["Settings"]
    for r in range(1, ws.max_row + 1):
        k = ws.cell(r, 2).value
        if isinstance(k, str) and k.startswith("settings."):
            out[k] = ws.cell(r, 4).value
    return out


def check_protocol(rep: Report, path: str, wf, wv):
    g = "PROTOCOL"
    rep.ok("Settings" in wf.sheetnames and "Index" in wf.sheetnames, g, "Settings and Index tabs present")
    if "Settings" in wf.sheetnames:
        st = wf["Settings"]
        hdr = [st.cell(1, c).value for c in (2, 3, 4, 5)]
        rep.ok(hdr == ["id", "Setting", "Value", "Description"], g,
               f"Settings header row reads id | Setting | Value | Description (got {hdr}) — the add-in matches 'id' and 'Value' literally")
        rep.ok(st["A1"].value is None, g, "Settings column A empty")
    s = settings_map(wf)
    rep.ok(str(s.get("settings.smartmodelSpec")) == "6.0", g, f"settings.smartmodelSpec = 6.0 (got {s.get('settings.smartmodelSpec')!r})")
    rep.ok(bool(s.get("settings.companyId")), g, f"settings.companyId set ({s.get('settings.companyId')!r}) — must be the Drivepoint tenant id")
    for k in ("settings.modelStartDate", "settings.historicalStartDate", "settings.lastDateActuals"):
        rep.ok(isinstance(s.get(k), (datetime, date)), g, f"{k} is a real Excel date ({s.get(k)!r})")
    rep.ok("Plan Settings" not in wf.sheetnames, g,
           "no 'Plan Settings' tab copied from the source (it pins another file's SharePoint id; the server writes a fresh one)", "WARN")
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        rep.ok("xl/webextensions/webextension1.xml" in names and b"webextensiontaskpanes" in z.read("_rels/.rels"), g,
               "Drivepoint add-in WebExtension part + root relationship present")
    return s


def check_calc(rep: Report, wf, wv, allow_uncalculated: bool = False):
    g = "CALC"
    errs, uncached = [], []
    for ws in wf.worksheets:
        vs = wv[ws.title]
        for row in ws.iter_rows():
            for c in row:
                v = vs.cell(c.row, c.column).value
                if isinstance(v, str) and v.startswith(ERRORS):
                    errs.append(f"{ws.title}!{c.coordinate}")
                if is_formula(c.value) and v is None:
                    uncached.append(f"{ws.title}!{c.coordinate}")
    rep.ok(not errs, g, f"no error cells ({len(errs)}: {errs[:6]})")
    rep.ok(not uncached, g, f"every formula has a cached value ({len(uncached)} without: {uncached[:6]}) — recalc before hand-off"
           + (" (allowed: open in Excel, Ctrl+Alt+F9, save — BEFORE uploading)" if allow_uncalculated and uncached else ""),
           "WARN" if allow_uncalculated else "FAIL")


def spine_of(wv_ws):
    cols, dates = [], []
    c = SPINE_START
    while True:
        d = as_date(wv_ws.cell(2, c).value)
        if d is None:
            break
        cols.append(c)
        dates.append(d)
        c += 1
    return cols, dates


def fill_spine_from_formulas(wf, wv, settings: dict) -> int:
    """Uncalculated workbook: evaluate the standard spine formulas (row 2 EOMONTH chain, row 3 Actual/Forecast IF)
    so structure can still be checked. Returns the number of cells filled."""
    filled = 0
    cache = {"historicalStartDate": settings.get("settings.historicalStartDate"),
             "modelStartDate": settings.get("settings.modelStartDate"), "lastDateActuals": settings.get("settings.lastDateActuals")}
    st = wf["Settings"] if "Settings" in wf.sheetnames else None
    by_cell = {}
    if st:
        for r in range(1, st.max_row + 1):
            v = st.cell(r, 4).value
            if isinstance(v, (datetime, date)):
                by_cell[f"$D${r}"] = v if isinstance(v, datetime) else datetime(v.year, v.month, v.day)

    def eom(d: datetime, k: int) -> datetime:
        y, m = d.year + (d.month - 1 + k) // 12, (d.month - 1 + k) % 12 + 1
        return datetime(y + (m == 12), m % 12 + 1, 1) - timedelta(days=1)
    for ws in wf.worksheets:
        vs = wv[ws.title]
        if vs["C2"].value != "End of Period":
            continue
        c = SPINE_START
        while True:
            f = ws.cell(2, c).value
            if not is_formula(f):
                break
            m1 = re.fullmatch(r"=EOMONTH\(Settings!(\$D\$\d+),(-?\d+)\)", f)
            m2 = re.fullmatch(r"=EOMONTH\(\$?([A-Z]{1,3})\$?2,(-?\d+)\)", f)
            if m1 and m1.group(1) in by_cell:
                d = eom(by_cell[m1.group(1)], int(m1.group(2)))
            elif m2 and isinstance(vs[f"{m2.group(1)}2"].value, datetime):
                d = eom(vs[f"{m2.group(1)}2"].value, int(m2.group(2)))
            else:
                break
            if vs.cell(2, c).value is None:
                vs.cell(2, c).value = d
                filled += 1
            f3 = ws.cell(3, c).value
            m3 = re.fullmatch(r'=IF\(\$?[A-Z]{1,3}\$?2<=Settings!(\$D\$\d+),"Actual","Forecast"\)', str(f3 or ""))
            if m3 and m3.group(1) in by_cell and vs.cell(3, c).value is None:
                vs.cell(3, c).value = "Actual" if d <= by_cell[m3.group(1)] else "Forecast"
                filled += 1
            c += 1
    return filled


def check_tab(rep: Report, name: str, wf, wv, settings: dict):
    ws, vs = wf[name], wv[name]
    g = f"{name[:9]}"
    cols, dates = spine_of(vs)
    rep.ok(vs["C2"].value == "End of Period" and vs["C3"].value == "Period Type", g, "C2 = 'End of Period', C3 = 'Period Type'")
    if not cols:
        rep.add("FAIL", g, "SPINE: no date in K2 — the date spine must start in column K")
        return
    last_col = cols[-1]
    contiguous = all((dates[i + 1].year * 12 + dates[i + 1].month) - (dates[i].year * 12 + dates[i].month) == 1
                     for i in range(len(dates) - 1))
    rep.ok(contiguous and (all(month_end(d) for d in dates) or all(d.day == 1 for d in dates)), g,
           f"SPINE: {len(dates)} contiguous months K2:{get_column_letter(last_col)}2 ({dates[0]:%b %Y} → {dates[-1]:%b %Y})")
    stray2 = [get_column_letter(c) for c in range(last_col + 1, ws.max_column + 1) if vs.cell(2, c).value not in (None, "")]
    rep.ok(not stray2, g, f"SPINE: nothing in row 2 after the last date ({stray2[:5]})")
    pt = [vs.cell(3, c).value for c in cols]
    stray3 = [get_column_letter(c) for c in range(last_col + 1, ws.max_column + 1) if vs.cell(3, c).value not in (None, "")]
    ok_pt = all(x in ("Actual", "Forecast") for x in pt)
    n_act = sum(1 for x in pt if x == "Actual")
    monotone = pt == ["Actual"] * n_act + ["Forecast"] * (len(pt) - n_act)
    rep.ok(ok_pt and monotone and not stray3, g,
           f"SPINE: row 3 is Actual ×{n_act} then Forecast ×{len(pt) - n_act}, nothing past the spine "
           f"(plan sync rejects a row-3 value without a row-2 date){' — stray: ' + ', '.join(stray3[:4]) if stray3 else ''}")
    lda = settings.get("settings.lastDateActuals")
    if isinstance(lda, (datetime, date)) and n_act:
        la = dates[n_act - 1]
        rep.ok((la.year, la.month) == (lda.year, lda.month), g,
               f"SPINE: last Actual month {la:%b %Y} = Settings lastDateActuals {lda:%b %Y}")
    # ---- HISTORY ----------------------------------------------------------------------------
    hsd = settings.get("settings.historicalStartDate")
    if isinstance(hsd, (datetime, date)):
        rep.ok((dates[0].year, dates[0].month) <= (hsd.year, hsd.month), g,
               f"HISTORY: spine starts {dates[0]:%b %Y}, at or before Settings historicalStartDate {hsd:%b %Y} — "
               "history must sit IN the time series, not in a block underneath")
    rep.ok(n_act > 0, g, f"HISTORY: {n_act} Actual month(s) on the spine", "WARN" if not hsd else "FAIL")
    stacked = []
    for r in range(5, ws.max_row + 1):
        vals = [vs.cell(r, c).value for c in range(3, last_col + 1)]
        mon = sum(1 for v in vals if isinstance(v, str) and v.strip().upper()[:3] in MONTHS and len(v.strip()) <= 9)
        flags = sum(1 for v in vals if isinstance(v, str) and v.strip().lower() in ("act", "for", "actual", "forecast", "bud", "budget"))
        dts = sum(1 for c in cols if isinstance(vs.cell(r, c).value, (datetime, date)))
        if mon >= 6 or flags >= 6 or dts >= 6:
            stacked.append(r)
    rep.ok(not stacked, g, f"HISTORY: no second month-header / Act-For row below row 4 ({stacked[:6]}) — "
                           "an LY or actuals block stacked under the budget belongs in the spine's history columns")
    right = sorted({get_column_letter(c) for r in range(5, ws.max_row + 1) for c in range(last_col + 1, ws.max_column + 1)
                    if is_num(ws.cell(r, c).value) or is_formula(ws.cell(r, c).value)})
    rep.ok(not right, g, f"HISTORY: no Total / LY / variance columns right of the spine ({right[:6]}) — put them on a summary tab")
    # ---- MARKERS ----------------------------------------------------------------------------
    fc_all = [c for c, p in zip(cols, pt) if p == "Forecast"]
    act_cols = [c for c, p in zip(cols, pt) if p == "Actual"]
    kd, kr, badA, orphanB = [], [], [], []
    for r in range(5, ws.max_row + 1):
        a, b = ws.cell(r, 1).value, ws.cell(r, 2).value
        if isinstance(a, str) and KD_TOKEN in a:
            kd.append(r)
        elif isinstance(a, str) and KR_TOKEN in a:
            kr.append(r)
        elif a not in (None, "", "•") and r > 16:
            badA.append(r)
        if isinstance(b, str) and b.strip() and r > 16 and not (isinstance(a, str) and (KD_TOKEN in a or KR_TOKEN in a)) \
                and not re.match(r"^(metadata|settings)(___|\.)", b):
            orphanB.append(r)
    rep.ok(not badA, g, f"MARKERS: column A holds only Key Driver / Key Result markers (other text on rows {badA[:6]})")
    rep.ok(not orphanB, g, f"MARKERS: column B ids only on marked rows (ids without a marker on rows {orphanB[:6]})", "WARN")
    marked = kd + kr
    # Budget columns = forecast months the Key Drivers actually drive (most KD cells are typed inputs).
    # Forecast months fed from an R-tab / seed (e.g. the customer's own current-year forecast) are
    # reported, but held to the history rules, not the driver rules.
    imported = lambda v: is_formula(v) and re.search(r"'?[RF]\s*-\s*[^!']*'?!", v)  # noqa: E731
    fc_cols, seeded = [], []
    for c in fc_all:
        inputs = sum(1 for r in kd if ws.cell(r, c).value is not None and not is_formula(ws.cell(r, c).value))
        pulls = sum(1 for r in marked if imported(ws.cell(r, c).value))
        driven = (inputs >= max(1, len(kd) // 2)) if kd else True
        (fc_cols if driven and pulls < max(1, len(marked) // 3) else seeded).append(c)
    if seeded:
        rep.add("PASS", g, f"SPINE: {len(seeded)} forecast month(s) fed from an R-tab/seed "
                           f"({get_column_letter(seeded[0])}:{get_column_letter(seeded[-1])}); "
                           f"{len(fc_cols)} budget month(s) driven by Key Drivers"
                           + (f" ({get_column_letter(fc_cols[0])}:{get_column_letter(fc_cols[-1])})" if fc_cols else ""))
    rep.ok(bool(fc_cols) or not kd, g, "MEANING: Key Drivers hold inputs in the budget months")
    names = [str(ws.cell(r, 3).value or "").strip() for r in marked]
    rep.ok(all(names), g, f"MARKERS: every marked row has its friendly name in column C (missing on "
                          f"{[r for r, n in zip(marked, names) if not n][:6]})")
    dup = sorted({n for n in names if n and names.count(n) > 1})
    rep.ok(not dup, g, f"MARKERS: friendly names of marked rows are unique ({dup[:4]})")
    ids = [ws.cell(r, 2).value for r in marked]
    rep.ok(all(ids) and len(set(ids)) == len(ids), g, f"MARKERS: {len(kd)} Key Drivers + {len(kr)} Key Results, unique ids in B")
    # ---- MEANING ----------------------------------------------------------------------------
    kd_formula, kd_default = [], []
    for r in kd:
        fs = [ws.cell(r, c).value for c in fc_cols]
        if any(is_formula(x) for x in fs):
            own = all(not is_formula(x) or set(re.findall(r"\$?[A-Z]{1,3}\$?(\d+)", x)) <= {str(r)} for x in fs)
            (kd_default if own else kd_formula).append(r)
        if any(x is None for x in fs):
            kd_formula.append(r)
    rep.ok(not kd_formula, g, f"MEANING: Key Drivers are inputs in every forecast month (formula/blank on rows "
                              f"{sorted(set(kd_formula))[:6]} — those are results or calcs, not drivers)")
    if kd_default:
        rep.add("WARN", g, f"MEANING: Key Drivers with a formula default referencing only their own row (e.g. = same month LY): {kd_default[:6]}")
    kr_const = [r for r in kr if any(ws.cell(r, c).value is not None and not is_formula(ws.cell(r, c).value) for c in fc_cols)]
    kr_blank = [r for r in kr if any(ws.cell(r, c).value is None for c in fc_cols)]
    if kr_blank:
        rep.add("WARN", g, f"MEANING: Key Results blank in some forecast months (rows {kr_blank[:6]}) — intended?")
    kr_pure = [r for r in kr if fc_cols and all(is_formula(ws.cell(r, c).value) and not re.search(r"[A-Z]{1,3}\$?\d+", ws.cell(r, c).value[1:])
                                    for c in fc_cols)]
    rep.ok(not kr_const and not kr_pure, g, f"MEANING: Key Results are formulas of other rows in every forecast month "
                                            f"(hard-keyed on rows {kr_const[:6]}, constant-only formulas on {kr_pure[:6]} — those are drivers)")
    series = {r: [vs.cell(r, c).value for c in fc_cols] for r in range(17, ws.max_row + 1)}
    series = {r: v for r, v in series.items() if all(is_num(x) for x in v) and len({round(x, 6) for x in v}) > 1}
    pasted = []
    for r in kd:
        a = series.get(r)
        if not a or any(is_formula(ws.cell(r, c).value) for c in fc_cols):
            continue
        if "%" in str(ws.cell(r, fc_cols[0]).number_format) or all(abs(x) < 5 for x in a):
            continue  # a % / factor driver is already the driver, not a pasted $ amount
        for r2, b in series.items():
            if r2 == r or any(x == 0 for x in b):
                continue
            ks = [x / y for x, y in zip(a, b)]
            if 1e-6 < abs(ks[0]) < 1 and all(abs(k - ks[0]) <= max(1e-9, abs(ks[0]) * 1e-7) for k in ks) and any(a):
                pasted.append(f"row {r} `{ws.cell(r, 3).value}` = row {r2} `{ws.cell(r2, 3).value}` × {ks[0]:.6g}")
                break
    rep.ok(not pasted, g, "MEANING: no $ Key Driver is an exact % of another row (pasted % × base — make the % the driver): "
                          + ("; ".join(pasted[:4]) if pasted else "none"))
    unmarked = [r for r in range(17, ws.max_row + 1) if r not in marked and ws.cell(r, 3).value
                and fc_cols and all(is_num(ws.cell(r, c).value) for c in fc_cols)
                and any(ws.cell(r, c).value for c in fc_cols)]
    if unmarked:
        rep.add("WARN", g, "MEANING: hard-keyed forecast rows with no marker (should any be Key Drivers?): "
                + ", ".join(f"{r} `{ws.cell(r, 3).value}`" for r in unmarked[:8]))
    no_hist = [r for r in marked if act_cols and not any(is_num(vs.cell(r, c).value) and vs.cell(r, c).value != 0 for c in act_cols)]
    if no_hist:
        rep.add("WARN", g, "HISTORY: marked rows with no values in the Actual months (fine for budget-only drivers such as "
                "growth factors; otherwise seed them): " + ", ".join(f"{r} `{ws.cell(r, 3).value}`" for r in no_hist[:8]))


def check_ties(rep: Report, ties_path: str, wv):
    spec = json.load(open(ties_path))
    cache: dict = {}
    total = bad = 0
    msgs = []
    for t in spec:
        sheet, key = t["model"].split("!", 1)
        ws = wv[sheet]
        row = int(key) if key.isdigit() else next((r for r in range(1, ws.max_row + 1) if ws.cell(r, 2).value == key), None)
        if row is None:
            bad += 1
            msgs.append(f"{t['model']} not found")
            continue
        c1, c2 = (column_index_from_string(x) for x in t["cols"].split(":"))
        model = [ws.cell(row, c).value for c in range(c1, c2 + 1)]
        src_path = t["source"]
        if src_path not in cache:
            cache[src_path] = openpyxl.load_workbook(src_path, data_only=True)
        ssheet, rng = t["range"].split("!", 1)
        src = [c.value for row_ in cache[src_path][ssheet][rng] for c in row_]
        tol = float(t.get("tol", 0.005))
        for i, (m, s) in enumerate(zip(model, src)):
            total += 1
            mv, sv = float(m or 0) if is_num(m) or m is None else m, float(s or 0) if is_num(s) or s is None else s
            if not (is_num(mv) and is_num(sv) and abs(mv - sv) <= max(tol, abs(sv) * 1e-9)) and mv != sv:
                bad += 1
                if len(msgs) < 6:
                    msgs.append(f"{t['model']} {get_column_letter(c1 + i)}: {m} vs {t['range']}[{i}] {s}")
    rep.ok(bad == 0, "TIES", f"{total} values tie to the source ({bad} off{': ' + '; '.join(msgs) if msgs else ''})")


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--ties")
    ap.add_argument("--strict", action="store_true")
    ap.add_argument("--allow-uncalculated", action="store_true",
                    help="no formula engine available (e.g. a chat sandbox): uncached formulas become a WARN; "
                         "the user must recalculate in Excel before upload")
    a = ap.parse_args()
    wf = openpyxl.load_workbook(a.xlsx)
    wv = openpyxl.load_workbook(a.xlsx, data_only=True)
    rep = Report()
    settings = check_protocol(rep, a.xlsx, wf, wv)
    check_calc(rep, wf, wv, a.allow_uncalculated)
    if a.allow_uncalculated:
        n = fill_spine_from_formulas(wf, wv, settings)
        if n:
            rep.add("WARN", "CALC", f"evaluated {n} spine cells from their formulas (uncalculated workbook); value checks "
                                    "(pasted %, history presence, ties) need a recalculated file")
    tabs = [n for n in wf.sheetnames if not IGNORE_TABS.search(n) and wv[n]["C2"].value == "End of Period"
            and as_date(wv[n].cell(2, SPINE_START).value)]
    rep.ok(bool(tabs), "SPINE", f"time-series tabs found: {tabs}")
    for n in tabs:
        check_tab(rep, n, wf, wv, settings)
    if a.ties:
        check_ties(rep, a.ties, wv)
    print(f"# validate_drivepointified — {a.xlsx.split('/')[-1]}\n")
    print(rep.render())
    c = rep.counts()
    print(f"\n{c['PASS']} pass · {c['WARN']} warn · {c['FAIL']} fail")
    return 1 if c["FAIL"] or (a.strict and c["WARN"]) else 0


if __name__ == "__main__":
    sys.exit(main())
