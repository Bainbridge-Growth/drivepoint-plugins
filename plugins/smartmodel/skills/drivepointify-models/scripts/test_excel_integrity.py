#!/usr/bin/env python3
"""Self-test for the Gate 0 checkers (lint_xlsx.py, find_circular.py, style_gaps.py) on synthetic
workbooks: each one passes a clean file and catches the defect it exists for.

    python3 test_excel_integrity.py -v
"""
from __future__ import annotations

import re
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

import openpyxl
from openpyxl.styles import PatternFill

HERE = Path(__file__).resolve().parent
BLUE = PatternFill("solid", fgColor="FF64B0FF")


def run(script: str, *args: str) -> subprocess.CompletedProcess:
    return subprocess.run([sys.executable, str(HERE / script), *args], capture_output=True, text=True)


def rewrite_sheet(path: Path, fn) -> None:
    """Apply fn to the first worksheet's XML text inside the package."""
    src = zipfile.ZipFile(path)
    items = [(i, src.read(i.filename)) for i in src.infolist()]
    src.close()
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for info, data in items:
            if info.filename == "xl/worksheets/sheet1.xml":
                data = fn(data.decode("utf-8")).encode("utf-8")
            z.writestr(info, data)


class Gate0(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())

    def book(self, name: str, build) -> Path:
        wb = openpyxl.Workbook()
        build(wb.active)
        p = self.tmp / name
        wb.save(p)
        return p

    # ---- lint_xlsx.py
    def test_lint_clean_and_duplicate_cell(self):
        def build(ws):
            ws["C5"] = "Revenue"
            ws["K5"] = 100
            ws["L5"] = "=K5*1.1"
        p = self.book("lint.xlsx", build)
        self.assertEqual(run("lint_xlsx.py", str(p)).returncode, 0)
        # a second record for K5, the way a cell move that does not merge writes it
        rewrite_sheet(p, lambda x: re.sub(r'(<c r="K5"[^>]*>.*?</c>)', r'\1<c r="K5"/>', x, count=1))
        r = run("lint_xlsx.py", str(p))
        self.assertEqual(r.returncode, 1, r.stdout)
        self.assertIn("duplicate-cell", r.stdout)

    def test_lint_missing_part(self):
        p = self.book("rel.xlsx", lambda ws: ws.__setitem__("A1", 1))
        src = zipfile.ZipFile(p)
        items = [(i, src.read(i.filename)) for i in src.infolist()]
        src.close()
        with zipfile.ZipFile(p, "w") as z:
            for info, data in items:
                if info.filename == "xl/_rels/workbook.xml.rels":
                    data = data.replace(b"</Relationships>", b'<Relationship Id="rIdX" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/calcChain" Target="calcChain.xml"/></Relationships>')
                z.writestr(info, data)
        r = run("lint_xlsx.py", str(p))
        self.assertEqual(r.returncode, 1)
        self.assertIn("relationship-to-missing-part", r.stdout)

    # ---- find_circular.py
    def test_summary_sumifs_over_own_row_is_circular(self):
        def build(ws):
            for i, c in enumerate("KLMN"):
                ws[f"{c}2"] = f"=DATE(2026,{i + 1},1)"
                ws[f"{c}5"] = 10 * (i + 1)
            ws["F5"] = '=SUMIFS(5:5,$2:$2,">="&DATE(2026,1,1))'
        p = self.book("circ.xlsx", build)
        r = run("find_circular.py", str(p))
        self.assertEqual(r.returncode, 1, r.stdout)

        def fixed(ws):
            build(ws)
            ws["F5"] = '=SUMIFS($K5:$N5,$K$2:$N$2,">="&DATE(2026,1,1))'
        self.assertEqual(run("find_circular.py", str(self.book("ok.xlsx", fixed))).returncode, 0)

    def test_untaken_if_branch_still_circular(self):
        def build(ws):
            ws["F20"] = "Formula"
            ws["K19"] = 3
            ws["K20"] = '=IF($F20="Formula",K19*30,K45)'   # K45 is never evaluated...
            ws["K28"] = "=K20*0.6"
            ws["K29"] = "=K20*0.4"
            ws["K45"] = "=SUM(K28:K44)"                     # ...but Excel still sees the loop
        r = run("find_circular.py", str(self.book("if.xlsx", build)))
        self.assertEqual(r.returncode, 1, r.stdout)
        self.assertIn("1 circular", r.stdout)

    # ---- style_gaps.py
    def test_band_gap_found_fixed_and_allowed(self):
        def build(ws, gap=True, divider=False):
            for c in range(1, 15):
                if gap and c in (8, 9, 10):
                    continue
                ws.cell(1, c).fill = BLUE
            if divider:
                ws.cell(1, 7).fill = PatternFill("solid", fgColor="FF000000")
        r = run("style_gaps.py", str(self.book("gap.xlsx", build)))
        self.assertEqual(r.returncode, 1, r.stdout)
        self.assertIn("H:J", r.stdout)
        self.assertEqual(run("style_gaps.py", str(self.book("nogap.xlsx", lambda ws: build(ws, gap=False)))).returncode, 0)
        div = self.book("div.xlsx", lambda ws: build(ws, gap=False, divider=True))
        self.assertEqual(run("style_gaps.py", str(div)).returncode, 1)
        self.assertEqual(run("style_gaps.py", str(div), "--allow", "Sheet!G").returncode, 0)


if __name__ == "__main__":
    unittest.main()
